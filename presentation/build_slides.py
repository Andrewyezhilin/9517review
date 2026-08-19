"""Build the individual-presentation deck for Zhilin Ye (Group 95, GSOE9011)."""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

NAVY = RGBColor(0x1A, 0x27, 0x44)
BLUE = RGBColor(0x2F, 0x6F, 0xB3)
GOLD = RGBColor(0xB8, 0x86, 0x0B)
RED = RGBColor(0xC0, 0x39, 0x2B)
GREEN = RGBColor(0x2E, 0x7D, 0x52)
GREY = RGBColor(0x5A, 0x64, 0x76)
LIGHT = RGBColor(0xF2, 0xF4, 0xF8)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)
FONT = "Calibri"

prs = Presentation()
prs.slide_width = SLIDE_W
prs.slide_height = SLIDE_H
BLANK = prs.slide_layouts[6]


def add_slide():
    return prs.slides.add_slide(BLANK)


def rect(slide, x, y, w, h, fill, line=None):
    from pptx.enum.shapes import MSO_SHAPE
    sp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    sp.fill.solid()
    sp.fill.fore_color.rgb = fill
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line
        sp.line.width = Pt(1)
    sp.shadow.inherit = False
    return sp


def textbox(slide, x, y, w, h, lines, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    """lines: list of (text, size, color, bold, italic, space_after)."""
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for i, (text, size, color, bold, italic, space_after) in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space_after)
        run = p.add_run()
        run.text = text
        f = run.font
        f.name = FONT
        f.size = Pt(size)
        f.color.rgb = color
        f.bold = bold
        f.italic = italic
    return tb


def bullets(slide, x, y, w, h, items, size=16, color=NAVY, space=8, lead_color=GOLD):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(space)
        # optional bold lead: item can be (lead, rest) or plain string
        if isinstance(item, tuple):
            lead, rest = item
            r1 = p.add_run(); r1.text = "\u25aa "
            r1.font.name = FONT; r1.font.size = Pt(size); r1.font.color.rgb = lead_color; r1.font.bold = True
            r2 = p.add_run(); r2.text = lead
            r2.font.name = FONT; r2.font.size = Pt(size); r2.font.color.rgb = color; r2.font.bold = True
            r3 = p.add_run(); r3.text = rest
            r3.font.name = FONT; r3.font.size = Pt(size); r3.font.color.rgb = color
        else:
            r1 = p.add_run(); r1.text = "\u25aa "
            r1.font.name = FONT; r1.font.size = Pt(size); r1.font.color.rgb = lead_color; r1.font.bold = True
            r2 = p.add_run(); r2.text = item
            r2.font.name = FONT; r2.font.size = Pt(size); r2.font.color.rgb = color
    return tb


def header(slide, title, kicker=None):
    rect(slide, 0, 0, SLIDE_W, Inches(1.02), NAVY)
    rect(slide, 0, Inches(1.02), SLIDE_W, Inches(0.06), GOLD)
    lines = []
    if kicker:
        lines.append((kicker, 12, GOLD, True, False, 2))
    lines.append((title, 26, WHITE, True, False, 0))
    textbox(slide, Inches(0.55), Inches(0.08), Inches(11.5), Inches(0.9), lines,
            anchor=MSO_ANCHOR.MIDDLE)


def footer(slide, n):
    textbox(slide, Inches(0.55), Inches(7.08), Inches(9.5), Inches(0.35),
            [("Zhilin Ye  z5719978  |  Group 95  |  GSOE9011", 10, GREY, False, False, 0)])
    textbox(slide, Inches(12.35), Inches(7.08), Inches(0.6), Inches(0.35),
            [(str(n), 10, GREY, False, False, 0)], align=PP_ALIGN.RIGHT)


def style_table(tbl, header_fill=NAVY, header_color=WHITE, body_size=13, header_size=13):
    from pptx.oxml.ns import qn
    for r, row in enumerate(tbl.rows):
        for cell in row.cells:
            cell.margin_left = Inches(0.08)
            cell.margin_right = Inches(0.08)
            cell.margin_top = Inches(0.04)
            cell.margin_bottom = Inches(0.04)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            for p in cell.text_frame.paragraphs:
                for run in p.runs:
                    run.font.name = FONT
                    run.font.size = Pt(header_size if r == 0 else body_size)
                    if r == 0:
                        run.font.bold = True
                        run.font.color.rgb = header_color
                    else:
                        run.font.color.rgb = NAVY
            if r == 0:
                cell.fill.solid(); cell.fill.fore_color.rgb = header_fill
            else:
                cell.fill.solid()
                cell.fill.fore_color.rgb = LIGHT if r % 2 == 1 else WHITE


# ---------------------------------------------------------------- Slide 1: title
s = add_slide()
rect(s, 0, 0, SLIDE_W, SLIDE_H, NAVY)
rect(s, 0, Inches(4.62), SLIDE_W, Inches(0.06), GOLD)
textbox(s, Inches(0.9), Inches(1.35), Inches(11.5), Inches(2.4), [
    ("GSOE9011  |  GROUP 95  |  INDIVIDUAL PRESENTATION", 15, GOLD, True, False, 18),
    ("Experimental Design & Feasibility", 44, WHITE, True, False, 4),
    ("for Context-Aware Non-Recurring Traffic Congestion Forecasting", 24, RGBColor(0xC9, 0xD4, 0xE6), False, False, 0),
])
textbox(s, Inches(0.9), Inches(4.95), Inches(11.5), Inches(1.6), [
    ("Zhilin Ye   z5719978", 22, WHITE, True, False, 6),
    ("My role: defining evaluation episodes, leakage controls, metrics & decision rules,", 15, RGBColor(0xC9, 0xD4, 0xE6), False, False, 0),
    ("ablations, sensitivity analysis and risk planning", 15, RGBColor(0xC9, 0xD4, 0xE6), False, False, 0),
])

# ------------------------------------------------- Slide 2: project in one minute
s = add_slide()
header(s, "The project in one minute", "CONTEXT")
textbox(s, Inches(0.55), Inches(1.35), Inches(12.2), Inches(0.75), [
    ("Can incidents, weather and event schedules make 30-minute traffic forecasts reliable during disruptions,", 17, NAVY, False, False, 2),
    ("when models trained on routine sensor patterns fail exactly when early warning matters most?", 17, NAVY, False, False, 0),
])
box = rect(s, Inches(0.55), Inches(2.3), Inches(12.2), Inches(1.95), LIGHT)
textbox(s, Inches(0.85), Inches(2.45), Inches(11.6), Inches(1.75), [
    ("RESEARCH QUESTION (pre-registered)", 12, GOLD, True, False, 6),
    ("Can a context-aware deep-learning model integrating traffic sensors, incidents, weather and major-event "
     "schedules reduce the 30-minute MAE gap between recurring and non-recurring congestion by at least 20% "
     "versus a matched traffic-history baseline, while keeping recurring-congestion MAE within a 1% "
     "non-inferiority margin?", 16, NAVY, False, False, 0),
])
textbox(s, Inches(0.55), Inches(4.55), Inches(12.2), Inches(0.5), [
    ("Greater Los Angeles freeway network  |  TraffiDent / PeMS 2022-2024  |  NOAA weather  |  venue event calendars", 14, GREY, False, False, 0),
])
box2 = rect(s, Inches(0.55), Inches(5.2), Inches(12.2), Inches(1.55), WHITE, line=GOLD)
textbox(s, Inches(0.85), Inches(5.38), Inches(11.6), Inches(1.3), [
    ("MY ROLE", 12, GOLD, True, False, 6),
    ("The research question is a falsifiable, quantified claim. My job was to design the experiment that makes "
     "this claim testable: fairly, reproducibly, and within a 12-month budget.", 16, NAVY, True, False, 0),
])
footer(s, 2)

# ------------------------------------------------- Slide 3: contribution overview
s = add_slide()
header(s, "My contribution: four building blocks", "OVERVIEW")
cards = [
    ("1. Evaluation episodes", "What transparently counts as recurring vs non-recurring congestion; "
     "exposure rules for incidents, weather and events.", "\u00a7 4.3"),
    ("2. Leakage controls", "Chronological 60/20/20 splits, 7-day embargoes, episode-safe boundaries, "
     "forecast-time availability rules.", "\u00a7 4.3, 4.5"),
    ("3. Metrics & decision rules", "MAE / WAPE choice, episode-block bootstrap, pre-registered 20% gap "
     "reduction and 1% non-inferiority criteria.", "\u00a7 2.4, 4.5"),
    ("4. Ablations, sensitivity & risk", "E0-E8 comparison set, negative controls, alignment sensitivity "
     "envelope, month-1 feasibility audit.", "\u00a7 4.5, 4.6"),
]
positions = [(0.55, 1.5), (6.85, 1.5), (0.55, 4.15), (6.85, 4.15)]
for (title, body, ref), (x, y) in zip(cards, positions):
    rect(s, Inches(x), Inches(y), Inches(5.95), Inches(2.35), LIGHT)
    rect(s, Inches(x), Inches(y), Inches(0.1), Inches(2.35), GOLD)
    textbox(s, Inches(x + 0.35), Inches(y + 0.18), Inches(5.4), Inches(2.0), [
        (title, 18, NAVY, True, False, 6),
        (body, 14.5, GREY, False, False, 6),
        ("Proposal " + ref, 12, GOLD, True, False, 0),
    ])
textbox(s, Inches(0.55), Inches(6.62), Inches(12.2), Inches(0.4), [
    ("Together these turn the group's model idea into an auditable experiment.", 14, GREY, False, True, 0),
])
footer(s, 3)

# ------------------------------------------------- Slide 4: episode definition
s = add_slide()
header(s, "Defining what counts as non-recurring congestion", "BLOCK 1  \u00b7  \u00a7 4.3")
s.shapes.add_picture("presentation/fig_episode_definition.png",
                     Inches(0.55), Inches(1.3), width=Inches(8.35))
bullets(s, Inches(9.1), Inches(1.45), Inches(3.75), Inches(5.3), [
    ("Problem: ", "a detector fault, an ordinary peak and a crash queue all look like \"low speed\"."),
    ("Every threshold from training data only: ", "free-flow = 85th pct of overnight obs; episode = < 70% "
     "free-flow for \u2265 15 min."),
    ("Non-recurring needs two signals: ", "\u2265 20% / 15 km/h below expected profile AND an independently "
     "recorded disruption."),
    ("Labels stratify evaluation only ", "- never training targets, so the model cannot learn the labels."),
], size=13.5, space=10)
textbox(s, Inches(0.55), Inches(5.35), Inches(8.35), Inches(1.3), [
    ("Answers tutor feedback #9: \"define how it will be labelled experimentally\".", 13, GOLD, True, True, 4),
    ("Ambiguous cases are excluded from the primary analysis and reported in sensitivity analysis.", 13, GREY, False, False, 0),
])
footer(s, 4)

# ------------------------------------------------- Slide 5: leakage + fairness
s = add_slide()
header(s, "Leakage controls and a fair comparison", "BLOCK 2  \u00b7  \u00a7 4.3, 4.5")
s.shapes.add_picture("presentation/fig_chronological_split.png",
                     Inches(0.55), Inches(1.3), width=Inches(9.2))
bullets(s, Inches(10.0), Inches(1.45), Inches(2.9), Inches(3.0), [
    ("Random splits leak: ", "two windows of the same crash land in train AND test."),
    ("Forecast-time rules: ", "no final incident duration, no future weather at origin t."),
], size=13, space=10)
rect(s, Inches(0.55), Inches(4.45), Inches(12.2), Inches(2.05), LIGHT)
textbox(s, Inches(0.85), Inches(4.6), Inches(11.6), Inches(1.85), [
    ("FAIRNESS BY CONSTRUCTION (with Zhenyu)", 12, GOLD, True, False, 6),
])
bullets(s, Inches(0.85), Inches(5.05), Inches(11.6), Inches(1.4), [
    ("Same backbone: ", "baseline and context model share the Graph WaveNet encoder, splits, target and tuning budget."),
    ("E8 capacity control: ", "a widened, parameter-matched traffic-only model - so a win cannot be explained by \"more parameters\"."),
], size=14.5, space=6)
footer(s, 5)

# ------------------------------------------------- Slide 6: decision criteria
s = add_slide()
header(s, "Pre-registered decision rules and statistics", "BLOCK 3  \u00b7  \u00a7 2.4, 4.5")
rows = [
    ("Condition (ALL three required)", "Criterion"),
    ("1. Reliability gap  G = MAE(non-rec) - MAE(rec)", "Reduced by \u2265 20% vs Graph WaveNet baseline"),
    ("2. Non-recurring improvement is real", "95% episode-block bootstrap CI excludes zero (10,000 resamples, 5 seeds)"),
    ("3. Routine traffic is protected", "Recurring MAE upper confidence bound within 1% of baseline"),
]
tbl_shape = s.shapes.add_table(4, 2, Inches(0.55), Inches(1.4), Inches(12.2), Inches(2.3))
tbl = tbl_shape.table
tbl.columns[0].width = Inches(5.7)
tbl.columns[1].width = Inches(6.5)
for r, (a, b) in enumerate(rows):
    tbl.cell(r, 0).text = a
    tbl.cell(r, 1).text = b
style_table(tbl, body_size=14, header_size=14)
bullets(s, Inches(0.55), Inches(4.0), Inches(12.2), Inches(1.5), [
    ("Why MAE? ", "Directly interpretable in km/h; RMSE is dominated by extreme errors; WAPE replaces MAPE, "
     "which explodes in stop-and-go traffic near zero speed."),
    ("Why episode-block bootstrap? ", "Overlapping 5-min windows are not independent - resampling whole episodes "
     "preserves within-episode dependence. Diebold-Mariano as secondary diagnostic; Holm correction across ablations."),
], size=14.5, space=8)
rect(s, Inches(0.55), Inches(5.65), Inches(12.2), Inches(1.1), WHITE, line=RED)
textbox(s, Inches(0.85), Inches(5.78), Inches(11.6), Inches(0.9), [
    ("BUILT-IN LIE DETECTOR (E7)", 12, RED, True, False, 4),
    ("Re-run the model with context shuffled within time-of-day. If shuffled context still \"helps\", "
     "the pipeline is lying to us.", 15, NAVY, True, False, 0),
])
footer(s, 6)

# ------------------------------------------------- Slide 7: feasibility & risk
s = add_slide()
header(s, "Feasibility first: risks, triggers, contingencies", "BLOCK 4  \u00b7  \u00a7 4.6")
textbox(s, Inches(0.55), Inches(1.35), Inches(12.2), Inches(0.65), [
    ("The biggest risk is not building the model - it is having too few well-aligned non-recurring episodes "
     "to say anything meaningful.", 16, NAVY, True, False, 0),
])
rows = [
    ("Risk", "Early trigger", "My contingency"),
    ("Too few non-recurring episodes", "< 200 test episodes or < 30 per source",
     "Month-1 audit BEFORE building anything; relax completeness gate to 85%, extend to next-largest connected component"),
    ("Event data incomplete", "Missing venue dates or unverifiable reuse terms",
     "Official public schedules only; demote event ablation to exploratory before test access"),
    ("Context model overfits", "Validation gain vanishes across seeds, or shuffled context helps",
     "Fall back to simple concatenation; strengthen regularisation; keep negative controls"),
    ("Compute exceeds budget", "One trial > 24 h or memory limit",
     "Sensor cap (300-800 stations), mixed precision, fixed trials, early stopping"),
]
tbl_shape = s.shapes.add_table(5, 3, Inches(0.55), Inches(2.15), Inches(12.2), Inches(3.4))
tbl = tbl_shape.table
tbl.columns[0].width = Inches(2.9)
tbl.columns[1].width = Inches(3.4)
tbl.columns[2].width = Inches(5.9)
for r, row in enumerate(rows):
    for c, val in enumerate(row):
        tbl.cell(r, c).text = val
style_table(tbl, body_size=12.5, header_size=13)
textbox(s, Inches(0.55), Inches(6.05), Inches(12.2), Inches(0.85), [
    ("Even if the 20% target is not met, the project still succeeds: the threshold is a decision rule, "
     "not a promised result.", 15, GOLD, True, True, 0),
])
footer(s, 7)

# ------------------------------------------------- Slide 8: reflection & close
s = add_slide()
rect(s, 0, 0, SLIDE_W, SLIDE_H, NAVY)
rect(s, 0, Inches(2.52), SLIDE_W, Inches(0.05), GOLD)
textbox(s, Inches(1.0), Inches(1.0), Inches(11.3), Inches(1.5), [
    ("\u201cA model idea becomes research only when the experiment around it is trustworthy.\u201d",
     28, WHITE, True, True, 0),
], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
bullets(s, Inches(1.6), Inches(3.1), Inches(10.2), Inches(2.6), [
    ("Connective tissue: ", "aligned data-quality assumptions with Xin and comparison fairness with Zhenyu, "
     "so one set of definitions runs through data, model and evaluation."),
    ("Feasibility as design: ", "episode-count audit, compute caps and pre-declared fallbacks are built into "
     "the plan, not patched on later."),
    ("The real product: ", "an auditable answer to when and where context data actually helps - "
     "my job was to make that answer worth believing."),
], size=17, color=RGBColor(0xE4, 0xE9, 0xF2), space=14, lead_color=GOLD)
textbox(s, Inches(1.0), Inches(6.3), Inches(11.3), Inches(0.6), [
    ("Thank you  |  Zhilin Ye  z5719978", 16, GOLD, True, False, 0),
], align=PP_ALIGN.CENTER)

prs.save("presentation/group95_zhilin_ye_individual_presentation.pptx")
print("deck saved:", len(prs.slides.slides if hasattr(prs.slides, 'slides') else prs.slides._sldIdLst), "slides")
