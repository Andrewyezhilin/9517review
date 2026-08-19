"""Render the elevator-pitch script as a printable Word document."""
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

NAVY = RGBColor(0x1A, 0x27, 0x44)
GOLD = RGBColor(0xB8, 0x86, 0x0B)
RED = RGBColor(0xC0, 0x39, 0x2B)
GREY = RGBColor(0x5A, 0x64, 0x76)

doc = Document()
style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(11)

for section in doc.sections:
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)


def heading(text, size=16, color=NAVY, space_before=14):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.font.bold = True
    r.font.color.rgb = color
    return p


def script(text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(text)
    r.font.size = Pt(12.5)
    return p


def note(text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run(text)
    r.font.size = Pt(10)
    r.font.italic = True
    r.font.color.rgb = GREY
    return p


def body(text, size=10.5, color=None, bold=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    r.font.size = Pt(size)
    if color:
        r.font.color.rgb = color
    r.font.bold = bold
    return p


# ------------------------------------------------------------------ title
t = doc.add_paragraph()
t.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = t.add_run("Elevator Pitch (Video Pitch) — Speech Script")
r.font.size = Pt(20)
r.font.bold = True
r.font.color.rgb = NAVY

sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = sub.add_run("Zhilin Ye  |  z5719978  |  Group 95  |  GSOE9011\n"
                "Context-Aware Deep Learning for Non-Recurring Traffic Congestion Forecasting")
r.font.size = Pt(11)
r.font.color.rgb = GREY

body("使用说明：录制视频时照着读下面的英文稿（2 分钟版，约 300 词）。括号中的中文是表演提示，"
     "不要读出来。作业硬性要求：必须用自己的声音朗读（禁止 TTS 配音）；不可用 AI 生成视频；"
     "录制时让摄像头画面叠加在幻灯片角落可拿出镜加分。", size=10.5, color=RED)

# ------------------------------------------------------------------ 2-minute script
heading("2 分钟主版（配合 5 页 PPT：group95_elevator_pitch.pptx）", 14, GOLD)

heading("Slide 1 — 开场钩子（约 15 秒）", 12)
script('Hi, I\'m Zhilin Ye from Group 95. Let me ask you something: what good is a traffic '
       'forecast that works perfectly every ordinary day — but fails the moment a crash happens? '
       'That\'s exactly the problem our research tackles.')
note("提示：第一句直视镜头，\u201cfails\u201d 加重语气。")

heading("Slide 2 — 问题：预测盲区（约 30 秒）", 12)
script('Modern deep-learning models learn from millions of sensor readings, but those readings '
       'mostly capture routine weekday patterns. So when a collision, a rain cell, or a stadium '
       'crowd suddenly changes the road, the model keeps saying "everything looks normal" — you '
       'can see that grey line on the chart. The red area is the blind spot: the biggest '
       'forecasting errors happen exactly when travellers, freight and emergency services need '
       'early warning the most.')
note("提示：说到 grey line 和 red area 时，用鼠标或手势指一下图。")

heading("Slide 3 — 我们的方案（约 35 秒）", 12)
script('Our idea is simple to state: give the model the context it cannot see. We take a strong '
       'graph neural network — Graph WaveNet — and feed it three extra sources: live incident '
       'reports, weather observations, and event schedules, which are known hours in advance. '
       'A gated-attention layer learns when each source matters. And crucially, the model only '
       'ever sees information available at forecast time — no future data, ever.')
script('We turned this into one falsifiable question: can context cut the 30-minute error gap '
       'between normal and disrupted traffic by at least twenty percent, without hurting everyday '
       'forecasts by more than one percent?')
note("提示：\u201ctwenty percent\u201d 和 \u201cone percent\u201d 两个数字放慢、咬清楚。")

heading("Slide 4 — 凭什么可信（约 25 秒）", 12)
script('And you can trust the answer, because the experiment is built to be fair. Three years of '
       'real Los Angeles freeway data. A matched baseline with the same backbone, the same '
       'leak-free chronological splits, the same tuning budget. Source-by-source ablations, plus '
       'a shuffled-context negative control — a built-in lie detector that exposes fake gains. '
       'And success is defined in advance, with pre-registered decision rules.')

heading("Slide 5 — 价值收尾（约 15 秒）", 12)
script('Whatever the outcome, transport agencies get an auditable answer to one practical '
       'question: is context data worth investing in? Reliable forecasts, when it matters most. '
       'Thank you.')
note("提示：最后一句放慢，看镜头微笑收尾，停顿一秒再停止录制。")

doc.add_page_break()

# ------------------------------------------------------------------ 90-second version
heading("90 秒精简版（时限更紧时使用，约 220 词）", 14, GOLD, space_before=0)
script('Hi, I\'m Zhilin Ye from Group 95. What good is a traffic forecast that works every '
       'ordinary day — but fails the moment a crash happens?')
script('Deep-learning traffic models learn routine weekday patterns from millions of sensor '
       'readings. So when a collision, a rain cell or a stadium crowd suddenly changes the road, '
       'the model keeps predicting "normal". That red area on the chart is the blind spot — the '
       'biggest errors, exactly when early warning matters most.')
script('Our answer: give the model the context it cannot see. We take a strong graph neural '
       'network and add live incident reports, weather, and event schedules, fused by gated '
       'attention — using only information available at forecast time.')
script('We test one falsifiable claim on three years of Los Angeles freeway data: context must '
       'cut the 30-minute error gap between normal and disrupted traffic by at least twenty '
       'percent, without hurting everyday forecasts by more than one percent. A matched baseline, '
       'leak-free splits, ablations and a shuffled-context negative control keep the comparison '
       'honest — and success is defined before we ever touch the test set.')
script('Whatever the outcome, transport agencies get an auditable answer: is context data worth '
       'investing in? Reliable forecasts, when it matters most. Thank you.')

# ------------------------------------------------------------------ recording guide
heading("录制指南（对照作业要求）", 14, GOLD)
body("1. 本人配音：直接对着幻灯片朗读稿子，不要用任何配音软件。先朗读两三遍熟悉节奏。", 11)
body("2. 视觉材料：使用 group95_elevator_pitch.pptx。建议自己微调一两处措辞，确保能对每页内容负责，"
     "并按课程 AI 使用政策做披露。", 11)
body("3. 出镜加分：推荐用 PowerPoint 的\u201c录制幻灯片演示\u201d功能，摄像头画面自动叠加在右下角，"
     "录完直接导出 MP4。或用 Zoom 共享屏幕 + 开摄像头 + 本地录制。", 11)
body("4. 细节：安静房间；麦克风离嘴 20\u201330 厘米；开场和收尾看镜头；某页说错只需重录该页"
     "（PowerPoint 按页保存录音）；先用手机计时朗读全稿确认不超时。", 11)
body("5. 上传：导出 1080p MP4 \u2192 上传 YouTube 选 Unlisted（或 OneDrive 开启链接共享）\u2192 "
     "把链接提交到 portal。提交前用无痕窗口验证链接无需登录即可播放。", 11)

doc.save("presentation/elevator_pitch_script.docx")
print("docx saved")
