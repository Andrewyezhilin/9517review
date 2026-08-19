"""Build the elevator-pitch (video pitch) deck covering the whole Group 95 proposal."""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

NAVY = RGBColor(0x1A, 0x27, 0x44)
BLUE = RGBColor(0x2F, 0x6F, 0xB3)
GOLD = RGBColor(0xB8, 0x86, 0x0B)
RED = RGBColor(0xC0, 0x39, 0x2B)
GREEN = RGBColor(0x2E, 0x7D, 0x52)
GREY = RGBColor(0x5A, 0x64, 0x76)
LIGHT = RGBColor(0xF2, 0xF4, 0xF8)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
SOFT = RGBColor(0xC9, 0xD4, 0xE6)

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)
FONT = "Calibri"

prs = Presentation()
prs.slide_width = SLIDE_W
prs.slide_height = SLIDE_H
BLANK = prs.slide_layouts[6]


def add_slide():
    return prs.slides.add_slide(BLANK)


def rect(slide, x, y, w, h, fill, line=None):
    sp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    sp.fill.solid()
    sp.fill.fore_color.rgb = fill
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line
        sp.line.width = Pt(1.25)
    sp.shadow.inherit = False
    return sp


def textbox(slide, x, y, w, h, lines, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for i, (text, size, color, bold, italic, space_after) in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space_after)
        run = p.add_run()
        run.text = text
        f = run.font
        f.name = FONT
        f.size = Pt(size)
        f.color.rgb = color
        f.bold = bold
        f.italic = italic
    return tb


def footer(slide, n):
    textbox(slide, Inches(0.55), Inches(7.1), Inches(10.0), Inches(0.32),
            [("Group 95 - Context-Aware Deep Learning for Non-Recurring Traffic Congestion Forecasting",
              10, GREY, False, False, 0)])
    textbox(slide, Inches(12.5), Inches(7.1), Inches(0.5), Inches(0.32),
            [(str(n), 10, GREY, False, False, 0)], align=PP_ALIGN.RIGHT)


# ---------------------------------------------------------------- Slide 1: hook title
s = add_slide()
rect(s, 0, 0, SLIDE_W, SLIDE_H, NAVY)
rect(s, 0, Inches(4.85), SLIDE_W, Inches(0.06), GOLD)
textbox(s, Inches(0.9), Inches(1.15), Inches(11.6), Inches(3.4), [
    ("GSOE9011  |  GROUP 95  |  ELEVATOR PITCH", 15, GOLD, True, False, 22),
    ("Traffic forecasts fail exactly when", 40, WHITE, True, False, 2),
    ("we need them most.", 40, WHITE, True, False, 16),
    ("Context-Aware Deep Learning for Non-Recurring Traffic Congestion Forecasting",
     20, SOFT, False, False, 0),
])
textbox(s, Inches(0.9), Inches(5.25), Inches(11.6), Inches(1.4), [
    ("Zhilin Ye   z5719978", 22, WHITE, True, False, 4),
    ("A 12-month research proposal - Greater Los Angeles freeway network", 15, SOFT, False, False, 0),
])

# ---------------------------------------------------------------- Slide 2: the blind spot
s = add_slide()
textbox(s, Inches(0.55), Inches(0.3), Inches(12.2), Inches(0.75), [
    ("The problem: a blind spot worth billions of lost hours", 30, NAVY, True, False, 0),
])
rect(s, Inches(0.55), Inches(1.05), Inches(3.2), Inches(0.05), GOLD)
s.shapes.add_picture("presentation/fig_pitch_blindspot.png",
                     Inches(0.85), Inches(1.35), width=Inches(11.6))
textbox(s, Inches(0.85), Inches(6.5), Inches(11.6), Inches(0.55), [
    ("Deep-learning models learn routine weekday patterns from millions of sensor readings - "
     "then a crash, a rain cell or a stadium crowd changes everything in minutes.",
     16, GREY, False, True, 0),
], align=PP_ALIGN.CENTER)
footer(s, 2)

# ---------------------------------------------------------------- Slide 3: our idea
s = add_slide()
textbox(s, Inches(0.55), Inches(0.3), Inches(12.2), Inches(0.75), [
    ("Our idea: give the model the context it cannot see", 30, NAVY, True, False, 0),
])
rect(s, Inches(0.55), Inches(1.05), Inches(3.2), Inches(0.05), GOLD)
s.shapes.add_picture("presentation/fig_pitch_architecture.png",
                     Inches(0.85), Inches(1.3), width=Inches(11.6))
rect(s, Inches(0.85), Inches(6.15), Inches(11.6), Inches(0.85), LIGHT)
textbox(s, Inches(1.15), Inches(6.27), Inches(11.0), Inches(0.65), [
    ("One falsifiable question: can context cut the 30-minute error gap between normal and disrupted "
     "traffic by at least 20% - without hurting everyday forecasts by more than 1%?",
     15.5, NAVY, True, False, 0),
])
footer(s, 3)

# ---------------------------------------------------------------- Slide 4: why trust it
s = add_slide()
textbox(s, Inches(0.55), Inches(0.3), Inches(12.2), Inches(0.75), [
    ("Why you can trust the answer", 30, NAVY, True, False, 0),
])
rect(s, Inches(0.55), Inches(1.05), Inches(3.2), Inches(0.05), GOLD)
cards = [
    ("REAL DATA, AT SCALE", BLUE,
     "3 years of 5-minute data (2022-2024), hundreds of LA freeway sensors, aligned incidents, "
     "NOAA weather, venue event calendars."),
    ("A FAIR FIGHT", GREEN,
     "Context model vs matched Graph WaveNet baseline: same backbone, same leak-free chronological "
     "splits, same tuning budget."),
    ("BUILT-IN LIE DETECTOR", RED,
     "Source-by-source ablations, a parameter-matched control, and a shuffled-context negative "
     "control that exposes fake gains."),
    ("SUCCESS DEFINED IN ADVANCE", GOLD,
     "Pre-registered decision rules with bootstrap confidence intervals - the experiment can pass, "
     "fail, or be inconclusive. All three are useful."),
]
positions = [(0.55, 1.45), (6.85, 1.45), (0.55, 4.2), (6.85, 4.2)]
for (title, color, body), (x, y) in zip(cards, positions):
    rect(s, Inches(x), Inches(y), Inches(5.95), Inches(2.45), LIGHT)
    rect(s, Inches(x), Inches(y), Inches(5.95), Inches(0.14), color)
    textbox(s, Inches(x + 0.3), Inches(y + 0.32), Inches(5.4), Inches(2.0), [
        (title, 16, color, True, False, 8),
        (body, 14.5, NAVY, False, False, 0),
    ])
footer(s, 4)

# ---------------------------------------------------------------- Slide 5: payoff / close
s = add_slide()
rect(s, 0, 0, SLIDE_W, SLIDE_H, NAVY)
rect(s, 0, Inches(2.32), SLIDE_W, Inches(0.05), GOLD)
textbox(s, Inches(1.0), Inches(0.85), Inches(11.3), Inches(1.5), [
    ("Reliable forecasts when it matters most.", 34, WHITE, True, False, 0),
], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
who = [
    ("Travellers & freight", "trustworthy travel times during disruption"),
    ("Emergency services", "earlier warning, faster response"),
    ("Traffic-management centres", "act 30 minutes ahead, not behind"),
    ("Transport agencies", "evidence on which data feeds are worth buying"),
]
y = 2.75
for lead, rest in who:
    textbox(s, Inches(2.2), Inches(y), Inches(9.0), Inches(0.55), [
        ("\u25aa  " + lead + " - " + rest, 18, SOFT, False, False, 0),
    ])
    tb = s.shapes[-1]
    y += 0.62
textbox(s, Inches(1.0), Inches(5.5), Inches(11.3), Inches(1.0), [
    ("Whatever the outcome, agencies get an auditable answer to one question:", 16, SOFT, False, True, 4),
    ("is context data worth investing in?  Thank you.", 22, GOLD, True, False, 0),
], align=PP_ALIGN.CENTER)

prs.save("presentation/group95_elevator_pitch.pptx")
print("pitch deck saved:", len(prs.slides._sldIdLst), "slides")
